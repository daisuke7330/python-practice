"""YouTube でキーワード検索し、動画のタイトルと URL を表示する

使い方:
    python search_youtube.py [--query キーワード] [--max 件数] [--order 並び順] [--excel]

キーワードを指定せずに実行すると「Python 入門」で検索する。
再生回数・高評価数・投稿日・チャンネル名もあわせて表示し、--excel で Excel に保存する。
upload_to_drive.py と同じ credentials.json を使い、ログイン情報は token_youtube.json に保存する。
初回はブラウザが開くので、Google アカウントでログインして許可する。
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from google.auth.exceptions import RefreshError
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

BASE_DIR = Path(__file__).resolve().parent
CREDENTIALS_FILE = BASE_DIR / "credentials.json"  # Google Cloud で作った OAuth クライアントの鍵
TOKEN_FILE = BASE_DIR / "token_youtube.json"  # YouTube 用のログイン保存ファイル（ほかの API とは権限が違うので分ける）

# YouTube の情報を見るだけの権限（動画の投稿や削除はできない）
SCOPES = ["https://www.googleapis.com/auth/youtube.readonly"]

VIDEO_URL = "https://www.youtube.com/watch?v={}"

# 並び順（--order）の指定と、画面に出す名前
ORDERS = {
    "relevance": "関連度順",
    "viewCount": "再生回数順",
    "date": "新しい順",
    "rating": "評価順",
}


def get_credentials():
    creds = None
    if TOKEN_FILE.exists():
        creds = Credentials.from_authorized_user_file(str(TOKEN_FILE), SCOPES)

    if creds and creds.valid:
        return creds

    if creds and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
        except RefreshError:
            creds = None  # 保存したログインが失効していたら、ログインし直す

    if not (creds and creds.valid):
        if not CREDENTIALS_FILE.exists():
            sys.exit(
                f"エラー: {CREDENTIALS_FILE.name} が見つかりません。\n"
                "Google Cloud Console で OAuth クライアント（デスクトップ アプリ）を作成し、\n"
                f"ダウンロードした JSON を {CREDENTIALS_FILE} に置いてください。"
            )
        flow = InstalledAppFlow.from_client_secrets_file(str(CREDENTIALS_FILE), SCOPES)
        creds = flow.run_local_server(port=0)  # ブラウザが開いてログイン・許可を求める

    TOKEN_FILE.write_text(creds.to_json(), encoding="utf-8")
    return creds


def search_videos(service, query, max_results, order):
    # type=video で、チャンネルや再生リストを除いて動画だけを探す
    res = service.search().list(
        q=query,
        part="snippet",
        type="video",
        maxResults=max_results,
        order=order,
        regionCode="JP",
        relevanceLanguage="ja",
    ).execute()
    return [item["id"]["videoId"] for item in res.get("items", [])]


def get_video_details(service, video_ids):
    # 検索結果には再生回数が入っていないので、videos.list で動画ごとの数字をまとめて取る
    res = service.videos().list(part="snippet,statistics", id=",".join(video_ids)).execute()
    details = {}
    for item in res.get("items", []):
        stats = item.get("statistics", {})
        details[item["id"]] = {
            "title": item["snippet"]["title"],
            "channel": item["snippet"]["channelTitle"],
            "published": item["snippet"]["publishedAt"][:10],
            "views": int(stats.get("viewCount", 0)),
            "likes": int(stats["likeCount"]) if "likeCount" in stats else None,  # 高評価数を非公開にしている動画もある
            "url": VIDEO_URL.format(item["id"]),
        }
    return [details[v] for v in video_ids if v in details]  # 検索の並び順を保つ


def save_excel(videos, query, order):
    path = BASE_DIR / f"youtube_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "検索結果"
    ws.append([f"検索キーワード：{query}（{ORDERS[order]}）"])
    ws.append(["No", "タイトル", "チャンネル", "投稿日", "再生回数", "高評価数", "URL"])
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")
    for i, v in enumerate(videos, 1):
        ws.append([i, v["title"], v["channel"], v["published"], v["views"], v["likes"], v["url"]])
        ws.cell(row=i + 2, column=7).hyperlink = v["url"]
    for row in ws.iter_rows(min_row=3, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    for col, width in zip("ABCDEFG", (5, 60, 28, 12, 12, 10, 45)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return path


def parse_args():
    parser = argparse.ArgumentParser(description="YouTube でキーワード検索し、動画のタイトルと URL を表示します。")
    parser.add_argument("--query", default="Python 入門", help="検索キーワード（省略時: Python 入門）")
    parser.add_argument("--max", type=int, default=5, choices=range(1, 51), metavar="1-50", help="表示する件数（省略時: 5）")
    parser.add_argument(
        "--order",
        default="relevance",
        choices=ORDERS,
        help="並び順（relevance=関連度 / viewCount=再生回数 / date=新しい順 / rating=評価。省略時: relevance）",
    )
    parser.add_argument("--excel", action="store_true", help="結果を Excel ファイルに保存する")
    return parser.parse_args()


def main():
    args = parse_args()
    service = build("youtube", "v3", credentials=get_credentials())

    try:
        video_ids = search_videos(service, args.query, args.max, args.order)
        videos = get_video_details(service, video_ids) if video_ids else []
    except HttpError as e:
        sys.exit(f"✗ 検索に失敗しました（{e.status_code}: {e.reason}）")

    if not videos:
        print(f"「{args.query}」に一致する動画は見つかりませんでした。")
        return

    print(f"✓ 「{args.query}」の検索結果（{ORDERS[args.order]}・{len(videos)} 件）")
    for i, v in enumerate(videos, 1):
        likes = f"{v['likes']:,}" if v["likes"] is not None else "非公開"
        print(f"{i}. {v['title']}")
        print(f"   {v['channel']} ／ {v['published']} ／ 再生 {v['views']:,} 回 ／ 高評価 {likes}")
        print(f"   {v['url']}")

    if args.excel:
        print(f"\n✓ Excel に保存しました: {save_excel(videos, args.query, args.order).name}")


if __name__ == "__main__":
    main()
